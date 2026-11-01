"""Case-change workbook -> INCH tests.

Fixtures are built from the shipped template via ``append_rows``, so a test
can never encode a column layout the template does not actually have.
"""

from __future__ import annotations

import re

import pytest

from tara_automation.casechange import read_workbook, write_template
from tara_automation.casechange.example import build_example_workbook
from tara_automation.casechange.inch import InchWriter
from tara_automation.casechange.schema import CHANGE_SHEETS, SHEETS
from tara_automation.casechange.workbook import append_rows

BASE = {"ScenarioID": "S1", "Name": "one", "Description": "first",
        "OutputFormat": "EPC", "OutputVersion": 21, "Solve": "Y",
        "NewBusStart": 990001, "Enabled": "Y"}


def _wb(tmp_path, scenarios=None, **sheets):
    p = tmp_path / "changes.xlsx"
    write_template(p)
    append_rows(p, "Scenarios", scenarios or [BASE])
    for name, rows in sheets.items():
        append_rows(p, name, rows)
    return p


def _render(cs, sid="S1"):
    return InchWriter(cs, cs.scenario(sid), "base.epc", "out.epc").render()


# ------------------------------------------------------------- template -----
def test_template_has_every_sheet(tmp_path):
    from openpyxl import load_workbook
    p = write_template(tmp_path / "t.xlsx")
    names = load_workbook(p).sheetnames
    for sh in SHEETS:
        assert sh.name in names
    assert "README" in names


def test_template_example_rows_are_ignored(tmp_path):
    """A freshly generated template must read back as empty, not as data."""
    p = write_template(tmp_path / "t.xlsx")
    cs = read_workbook(p)
    assert cs.scenarios == []
    for sh in CHANGE_SHEETS:
        assert cs.rows[sh.name] == []


def test_append_rows_rejects_unknown_column(tmp_path):
    p = write_template(tmp_path / "t.xlsx")
    with pytest.raises(KeyError, match="Nonsense"):
        append_rows(p, "New_Buses", [{"Nonsense": 1}])


# ------------------------------------------------------------- coercion -----
def test_defaults_fill_blank_cells(tmp_path):
    p = _wb(tmp_path, New_Buses=[
        {"Scenarios": "ALL", "BusName": "NEWBUS", "BusNomVolt": 500,
         "AreaNum": 14}])
    cs = read_workbook(p)
    assert cs.ok, cs.problems
    row = cs.rows["New_Buses"][0]
    assert row["BusType"] == 1          # schema default, not blank
    assert row["BusVoltPU"] == 1.0


def test_circuit_id_survives_excel_numeric_coercion(tmp_path):
    p = _wb(tmp_path, New_Branches=[
        {"Scenarios": "ALL", "BusNumFr": 1, "BusNumTo": 2, "CKT": 1,
         "LineR": 0.01, "LineX": 0.1, "RateA": 100}])
    cs = read_workbook(p)
    assert cs.rows["New_Branches"][0]["CKT"] == "1"


def test_bad_choice_is_reported(tmp_path):
    p = _wb(tmp_path, Gen_Load_Scaling=[
        {"Scenarios": "ALL", "Type": "WIND", "SelectBy": "AREA",
         "Selector": "14", "Value": 10}])
    cs = read_workbook(p)
    assert any(p_.column == "Type" for p_ in cs.problems)


# ----------------------------------------------------------- validation -----
def test_unknown_scenario_tag_is_reported(tmp_path):
    p = _wb(tmp_path, New_Buses=[
        {"Scenarios": "S9", "BusName": "B", "BusNomVolt": 500,
         "AreaNum": 14}])
    cs = read_workbook(p)
    assert any("unknown scenario" in str(x) for x in cs.problems)


def test_duplicate_scenario_id_is_reported(tmp_path):
    p = _wb(tmp_path, scenarios=[BASE, dict(BASE, Name="dup")])
    cs = read_workbook(p)
    assert any("duplicate scenario" in str(x) for x in cs.problems)


def test_unresolved_bus_reference_is_reported(tmp_path):
    p = _wb(tmp_path, New_Generators=[
        {"Scenarios": "ALL", "BusNum": "@NOSUCH", "GenID": "1",
         "GenMW": 10, "GenMWMax": 10}])
    cs = read_workbook(p)
    assert any("@NOSUCH" in str(x) for x in cs.problems)


def test_zero_rating_branch_is_reported(tmp_path):
    p = _wb(tmp_path, New_Branches=[
        {"Scenarios": "ALL", "BusNumFr": 1, "BusNumTo": 2, "CKT": "1",
         "LineR": 0.01, "LineX": 0.1, "RateA": 0}])
    cs = read_workbook(p)
    assert any("unrated" in str(x) for x in cs.problems)


def test_duplicate_new_bus_number_is_reported(tmp_path):
    p = _wb(tmp_path, New_Buses=[
        {"Scenarios": "ALL", "BusNum": 990001, "BusName": "A",
         "BusNomVolt": 500, "AreaNum": 14},
        {"Scenarios": "ALL", "BusNum": 990001, "BusName": "B",
         "BusNomVolt": 500, "AreaNum": 14}])
    cs = read_workbook(p)
    assert any("used twice" in str(x) for x in cs.problems)


# ------------------------------------------------------ bus assignment ------
def test_blank_bus_numbers_are_assigned_and_referenced(tmp_path):
    p = _wb(tmp_path,
            New_Buses=[{"Scenarios": "ALL", "BusName": "POI",
                        "BusNomVolt": 500, "AreaNum": 14}],
            New_Generators=[{"Scenarios": "ALL", "BusNum": "@POI",
                             "GenID": "1", "GenMW": 100, "GenMWMax": 100}])
    cs = read_workbook(p)
    assert cs.ok, cs.problems
    assert cs.scenario("S1").bus_names == {"POI": 990001}
    text = _render(cs)
    assert "990001,POI,500" in text
    assert re.search(r"#ADD_GEN .*\n990001,1,", text)


def test_explicit_bus_number_is_not_reassigned(tmp_path):
    p = _wb(tmp_path, New_Buses=[
        {"Scenarios": "ALL", "BusNum": 777001, "BusName": "POI",
         "BusNomVolt": 500, "AreaNum": 14},
        {"Scenarios": "ALL", "BusName": "COLL", "BusNomVolt": 34.5,
         "AreaNum": 14}])
    cs = read_workbook(p)
    assert cs.scenario("S1").bus_names == {"POI": 777001, "COLL": 990001}


# ------------------------------------------------------ scenario tagging ----
def test_rows_are_partitioned_by_scenario_tag(tmp_path):
    scen = [BASE, dict(BASE, ScenarioID="S2", Name="two")]
    p = _wb(tmp_path, scenarios=scen, New_Buses=[
        {"Scenarios": "S1", "BusName": "ONLY1", "BusNomVolt": 500,
         "AreaNum": 14},
        {"Scenarios": "ALL", "BusName": "BOTH", "BusNomVolt": 500,
         "AreaNum": 14}])
    cs = read_workbook(p)
    assert {r["BusName"] for r in cs.rows_for("New_Buses", "S1")} == {
        "ONLY1", "BOTH"}
    assert {r["BusName"] for r in cs.rows_for("New_Buses", "S2")} == {"BOTH"}


def test_disabled_scenario_is_excluded(tmp_path):
    scen = [BASE, dict(BASE, ScenarioID="S2", Name="two", Enabled="N")]
    p = _wb(tmp_path, scenarios=scen)
    cs = read_workbook(p)
    assert [s.id for s in cs.enabled_scenarios()] == ["S1"]


# --------------------------------------------------------------- INCH -------
def test_section_order_puts_buses_before_branches(tmp_path):
    p = _wb(tmp_path,
            New_Buses=[{"Scenarios": "ALL", "BusName": "POI",
                        "BusNomVolt": 500, "AreaNum": 14}],
            New_Branches=[{"Scenarios": "ALL", "BusNumFr": "@POI",
                           "BusNumTo": 14005, "CKT": "1", "LineR": 0.001,
                           "LineX": 0.01, "RateA": 1000}])
    text = _render(read_workbook(p))
    assert text.index("#ADD_BUS") < text.index("#ADD_BRANCH")
    assert text.index("#ADD_BRANCH") < text.index("#SAVE")


def test_modify_records_carry_only_filled_columns(tmp_path):
    p = _wb(tmp_path, Upgrades=[
        {"Scenarios": "ALL", "Element": "BRANCH", "BusNumFr": 1,
         "BusNumTo": 2, "CKT": "1", "RateA": 1200}])
    text = _render(read_workbook(p))
    assert "#MODIFY_BRANCH [BusNumFr,BusNumTo,CKT,RateA]" in text
    assert "LineX" not in text          # impedance must not be disturbed


def test_modify_rows_group_by_column_set(tmp_path):
    """Rows changing different fields need separate sections — one header
    cannot describe two different column lists."""
    p = _wb(tmp_path, Upgrades=[
        {"Scenarios": "ALL", "Element": "BRANCH", "BusNumFr": 1,
         "BusNumTo": 2, "CKT": "1", "RateA": 1200},
        {"Scenarios": "ALL", "Element": "BRANCH", "BusNumFr": 3,
         "BusNumTo": 4, "CKT": "1", "Status": 0}])
    text = _render(read_workbook(p))
    assert text.count("#MODIFY_BRANCH") == 2


def test_empty_columns_are_dropped_from_add_records(tmp_path):
    p = _wb(tmp_path, New_Buses=[
        {"Scenarios": "ALL", "BusName": "POI", "BusNomVolt": 500,
         "AreaNum": 14}])                      # ZoneNum / OwnerNum blank
    text = _render(read_workbook(p))
    header = next(l for l in text.splitlines() if l.startswith("#ADD_BUS"))
    assert "ZoneNum" not in header and "OwnerNum" not in header
    assert ",," not in text                    # no empty fields anywhere


def test_generator_defaults_are_derived(tmp_path):
    p = _wb(tmp_path, New_Generators=[
        {"Scenarios": "ALL", "BusNum": 14005, "GenID": "1", "GenMW": 200,
         "GenMWMax": 300}])
    text = _render(read_workbook(p))
    row = next(l for l in text.splitlines() if l.startswith("14005,"))
    assert "98.58" in row and "-98.58" in row   # 0.3286 x Pmax, mirrored
    assert row.split(",")[12] == "300"          # MBase defaults to Pmax


def test_names_with_spaces_are_quoted(tmp_path):
    p = _wb(tmp_path, New_Branches=[
        {"Scenarios": "ALL", "BusNumFr": 1, "BusNumTo": 2, "CKT": "1",
         "LineR": 0.001, "LineX": 0.01, "RateA": 1000,
         "BranchName": "EAGLE EYE TIE"}])
    assert '"EAGLE EYE TIE"' in _render(read_workbook(p))


def test_rate_b_and_c_inherit(tmp_path):
    p = _wb(tmp_path, New_Branches=[
        {"Scenarios": "ALL", "BusNumFr": 1, "BusNumTo": 2, "CKT": "1",
         "LineR": 0.001, "LineX": 0.01, "RateA": 1000}])
    text = _render(read_workbook(p))
    row = next(l for l in text.splitlines() if l.startswith("1,2,1,"))
    assert row.split(",")[7:10] == ["1000", "1000", "1000"]


def test_scaling_apply_order_controls_placement(tmp_path):
    rows = [{"Scenarios": "ALL", "Type": "LOAD", "SelectBy": "AREA",
             "Selector": "14", "Value": 250, "ApplyOrder": "BEFORE_ADDS"}]
    p = _wb(tmp_path, Gen_Load_Scaling=rows, New_Buses=[
        {"Scenarios": "ALL", "BusName": "POI", "BusNomVolt": 500,
         "AreaNum": 14}])
    text = _render(read_workbook(p))
    assert text.index("#SCALE_LOAD") < text.index("#ADD_BUS")


def test_subsystem_selector_is_not_numeric(tmp_path):
    p = _wb(tmp_path, Gen_Load_Scaling=[
        {"Scenarios": "ALL", "Type": "GEN", "SelectBy": "SUBSYSTEM",
         "Selector": "AZ", "Value": 100}])
    text = _render(read_workbook(p))
    assert "#SCALE_GEN [SubSysName," in text
    assert "AZ" in text


def test_solve_can_be_skipped(tmp_path):
    p = _wb(tmp_path, scenarios=[dict(BASE, Solve="N")])
    text = _render(read_workbook(p))
    assert "#SOLVE_FDEC" not in text
    assert "#SAVE" in text


def test_save_uses_scenario_format_and_version(tmp_path):
    p = _wb(tmp_path, scenarios=[dict(BASE, OutputFormat="RAW",
                                      OutputVersion=33)])
    cs = read_workbook(p)
    text = InchWriter(cs, cs.scenario("S1"), "b.epc", r"C:\out\new.raw").render()
    assert "#SAVE [Format,Version,FileName]" in text
    assert r"RAW,33,C:\out\new.raw" in text


def test_bad_output_version_is_reported(tmp_path):
    p = _wb(tmp_path, scenarios=[dict(BASE, OutputFormat="EPC",
                                      OutputVersion=22)])
    cs = read_workbook(p)
    assert any("EPC 18, 19 or 21" in str(x) for x in cs.problems)


def test_section_header_is_one_line(tmp_path):
    """Manual 6.1: the keyword and its bracketed column list must not wrap."""
    p = tmp_path / "ex.xlsx"
    build_example_workbook(p)
    cs = read_workbook(p)
    for sc in cs.enabled_scenarios():
        for line in _render(cs, sc.id).splitlines():
            if line.startswith("#") and "[" in line:
                assert line.rstrip().endswith("]"), line


# -------------------------------------------------------------- example -----
def test_example_workbook_is_valid_and_renders(tmp_path):
    p = tmp_path / "ex.xlsx"
    build_example_workbook(p)
    cs = read_workbook(p)
    assert cs.ok, [str(x) for x in cs.problems]
    assert [s.id for s in cs.enabled_scenarios()] == ["S1", "S2", "S3", "S4"]

    families = {"S1": ["#ADD_BUS", "#ADD_GEN", "#ADD_LOAD", "#ADD_BRANCH",
                       "#ADD_TRANSFORMER"],
                "S2": ["#MODIFY_GEN", "#MODIFY_BRANCH"],
                "S3": ["#SCALE_LOAD", "#SCALE_GEN"],
                "S4": ["#TAP"]}
    for sid, expected in families.items():
        text = _render(cs, sid)
        for keyword in expected:
            assert keyword in text, f"{sid} missing {keyword}"
        assert "#SAVE" in text


def test_example_tap_bus_is_created_before_it_is_used(tmp_path):
    p = tmp_path / "ex.xlsx"
    build_example_workbook(p)
    cs = read_workbook(p)
    text = _render(cs, "S4")
    assert text.index("#TAP") < text.index("#ADD_TRANSFORMER")
    tap_bus = cs.scenario("S4").bus_names["EETAP500"]
    xfmr = next(l for l in text.splitlines()
                if l.startswith(f"{tap_bus},"))
    assert xfmr.split(",")[0] == str(tap_bus)


# ------------------------------------------------- engine-side parsing ------
WARN_SUM = """LICENSEE  Warnings Summary[TARA Ver 2501.1.0.0 64-bit]

 Loadflow Case: 30HS_22-G7_pre.epc

 Severity Level: 1-info; 2-error; 3-potentially critical; 4-critical error,
WarnID, #Warn, Sever,Message,File,
  1100,     5,0-comment  , Number of completed input files,
  1200,     8,1-info     , Float Value out of range, File 30HS_22-G7_pre.epc
  1001,     1,2-error !  , Keyword [pvalueoption] not found, File S3.inch
  1607,     1,2-error !  , Can't add second transformer, the same buses,
  1652,     1,2-error !  , Can't find regulated bus for Transformer, File S3.inch
"""


def test_parse_warn_sum_reads_ids_and_files(tmp_path):
    from tara_automation.commfile.casechange import parse_warn_sum
    p = tmp_path / "Warn_Sum.csv"
    p.write_text(WARN_SUM, encoding="utf-8")
    rows = parse_warn_sum(p)
    by_id = {r["id"]: r for r in rows}
    assert set(by_id) == {1100, 1200, 1001, 1607, 1652}
    assert by_id[1001]["file"] == "S3.inch"
    assert by_id[1200]["case_noise"] is True
    assert by_id[1652]["case_noise"] is False
    # A message containing a comma must survive intact.
    assert by_id[1607]["message"] == "Can't add second transformer, the same buses"


def test_parse_warn_sum_missing_file_is_empty(tmp_path):
    from tara_automation.commfile.casechange import parse_warn_sum
    assert parse_warn_sum(tmp_path / "nope.csv") == []


def test_inch_errors_fail_a_run_despite_exit_zero(tmp_path):
    """The SCALE_LOAD misfire exited 0 and wrote a case. It must still fail."""
    from tara_automation.commfile.casechange import CaseChangeEngine
    from tara_automation.run_manager import Run

    case = tmp_path / "out.epc"
    case.write_bytes(b"x" * 50_000)
    run = Run(number=1, root=tmp_path)
    for d in (run.comm_dir, run.results_dir, run.logs_dir):
        d.mkdir(parents=True, exist_ok=True)
    (run.results_dir / "Warn_Sum.csv").write_text(WARN_SUM, encoding="utf-8")
    run.manifest = {"casechange": {"out_case": str(case)}, "tara_exit_code": 0}

    checks = CaseChangeEngine.check_outputs(run)
    assert checks["case_plausible"] is True
    assert checks["ok"] is False
    assert {e["id"] for e in checks["inch_errors"]} == {1001, 1607, 1652}
    assert "did not apply as written" in CaseChangeEngine.summary_line(checks)


def test_clean_run_passes(tmp_path):
    from tara_automation.commfile.casechange import CaseChangeEngine
    from tara_automation.run_manager import Run

    case = tmp_path / "out.epc"
    case.write_bytes(b"x" * 50_000)
    run = Run(number=1, root=tmp_path)
    for d in (run.comm_dir, run.results_dir, run.logs_dir):
        d.mkdir(parents=True, exist_ok=True)
    (run.results_dir / "Warn_Sum.csv").write_text(
        "WarnID, #Warn, Sever,Message,File,\n"
        "  1200,     8,1-info     , Float Value out of range, File case.epc\n",
        encoding="utf-8")
    run.manifest = {"casechange": {"out_case": str(case)}, "tara_exit_code": 0}
    assert CaseChangeEngine.check_outputs(run)["ok"] is True


def test_truncated_case_fails(tmp_path):
    from tara_automation.commfile.casechange import CaseChangeEngine
    from tara_automation.run_manager import Run

    case = tmp_path / "out.epc"
    case.write_bytes(b"")               # the 0-byte RAWD trap, generalised
    run = Run(number=1, root=tmp_path)
    for d in (run.comm_dir, run.results_dir, run.logs_dir):
        d.mkdir(parents=True, exist_ok=True)
    run.manifest = {"casechange": {"out_case": str(case)}, "tara_exit_code": 0}
    checks = CaseChangeEngine.check_outputs(run)
    assert checks["ok"] is False
    assert "bytes" in CaseChangeEngine.summary_line(checks)


def test_read_save_target_round_trips(tmp_path):
    from tara_automation.casechange.inch import read_save_target
    p = _wb(tmp_path, New_Buses=[
        {"Scenarios": "ALL", "BusName": "POI", "BusNomVolt": 500,
         "AreaNum": 14}])
    cs = read_workbook(p)
    inch = tmp_path / "s1.inch"
    InchWriter(cs, cs.scenario("S1"), "base.epc",
               tmp_path / "new.epc").write(inch)
    fmt, ver, out = read_save_target(inch)
    assert (fmt, ver) == ("EPC", 21)
    assert out == tmp_path / "new.epc"


def test_read_save_target_none_when_absent(tmp_path):
    from tara_automation.casechange.inch import read_save_target
    p = tmp_path / "x.inch"
    p.write_text("#ADD_BUS [BusNum]\n1\n#END\n", encoding="utf-8")
    assert read_save_target(p) is None


def test_transformer_carries_control_bus(tmp_path):
    """Run 32: W1CtrlBus is the ONE required control column — a record with
    only W1AdjStatus was rejected, one with only W1CtrlBus loaded. 0 is a
    valid control bus and is the fixed-tap case."""
    p = _wb(tmp_path, New_Transformers=[
        {"Scenarios": "ALL", "W1Bus": 990001, "W2Bus": 990002, "CKT": "1",
         "X1_2": 0.03, "W1RateA": 350}])
    text = _render(read_workbook(p))
    header = next(l for l in text.splitlines()
                  if l.startswith("#ADD_TRANSFORMER"))
    cols = header[header.index("[") + 1:header.index("]")].split(",")
    assert "W1CtrlBus" in cols
    # Ignored by TARA on an ADD — emitting them is noise that implies control
    # we do not have.
    assert "W1AdjStatus" not in cols and "W1ANG" not in cols
    values = next(l for l in text.splitlines()
                  if l.startswith("990001,990002,")).split(",")
    assert values[cols.index("W1CtrlBus")] == "0"


def test_transformer_gets_a_sane_tap_band(tmp_path):
    """Run 33: with no target band TARA stores +/-99999 on every added
    transformer, since they are all created as tap-controlling."""
    p = _wb(tmp_path, New_Transformers=[
        {"Scenarios": "ALL", "W1Bus": 990001, "W2Bus": 990002, "CKT": "1",
         "X1_2": 0.03, "W1RateA": 350}])
    text = _render(read_workbook(p))
    header = next(l for l in text.splitlines()
                  if l.startswith("#ADD_TRANSFORMER"))
    cols = header[header.index("[") + 1:header.index("]")].split(",")
    values = next(l for l in text.splitlines()
                  if l.startswith("990001,990002,")).split(",")
    assert values[cols.index("W1TargetMax")] == "1.05"
    assert values[cols.index("W1TargetMin")] == "0.95"


def test_fourth_rating_slot_is_filled(tmp_path):
    """Run 33: TARA leaves rate slot 4 at 0 unless RateD / W1Rate4 is given,
    which would make the element unrated for a study selecting rating 4."""
    p = _wb(tmp_path,
            New_Branches=[{"Scenarios": "ALL", "BusNumFr": 1, "BusNumTo": 2,
                           "CKT": "1", "LineR": 0.001, "LineX": 0.01,
                           "RateA": 1000}],
            New_Transformers=[{"Scenarios": "ALL", "W1Bus": 990001,
                               "W2Bus": 990002, "CKT": "1", "X1_2": 0.03,
                               "W1RateA": 350}])
    text = _render(read_workbook(p))
    bh = next(l for l in text.splitlines() if l.startswith("#ADD_BRANCH"))
    bcols = bh[bh.index("[") + 1:bh.index("]")].split(",")
    brow = next(l for l in text.splitlines() if l.startswith("1,2,1,")).split(",")
    assert brow[bcols.index("RateD")] == "1000"
    th = next(l for l in text.splitlines() if l.startswith("#ADD_TRANSFORMER"))
    tcols = th[th.index("[") + 1:th.index("]")].split(",")
    trow = next(l for l in text.splitlines()
                if l.startswith("990001,990002,")).split(",")
    assert trow[tcols.index("W1Rate4")] == "350"


def test_nameplate_percent_impedance_is_converted(tmp_path):
    """Run 32: TARA ignores CZ and always reads R/X as system-base per unit,
    so a nameplate percentage must be converted before it is written.
    10.5% on 350 MVA -> 0.105 x 100/350 = 0.03."""
    p = _wb(tmp_path, New_Transformers=[
        {"Scenarios": "ALL", "W1Bus": 990001, "W2Bus": 990002, "CKT": "1",
         "ImpedanceBasis": "NAMEPLATE_PCT", "X1_2": 10.5, "SBAS1": 350,
         "W1RateA": 350}])
    text = _render(read_workbook(p))
    header = next(l for l in text.splitlines()
                  if l.startswith("#ADD_TRANSFORMER"))
    cols = header[header.index("[") + 1:header.index("]")].split(",")
    values = next(l for l in text.splitlines()
                  if l.startswith("990001,990002,")).split(",")
    assert values[cols.index("X1-2")] == "0.03"


def test_system_basis_impedance_is_passed_through(tmp_path):
    p = _wb(tmp_path, New_Transformers=[
        {"Scenarios": "ALL", "W1Bus": 990001, "W2Bus": 990002, "CKT": "1",
         "X1_2": 0.03, "SBAS1": 350, "W1RateA": 350}])
    text = _render(read_workbook(p))
    header = next(l for l in text.splitlines()
                  if l.startswith("#ADD_TRANSFORMER"))
    cols = header[header.index("[") + 1:header.index("]")].split(",")
    values = next(l for l in text.splitlines()
                  if l.startswith("990001,990002,")).split(",")
    assert values[cols.index("X1-2")] == "0.03"


def test_nameplate_percent_without_a_base_is_reported(tmp_path):
    p = _wb(tmp_path, New_Transformers=[
        {"Scenarios": "ALL", "W1Bus": 990001, "W2Bus": 990002, "CKT": "1",
         "ImpedanceBasis": "NAMEPLATE_PCT", "X1_2": 10.5, "W1RateA": 0}])
    cs = read_workbook(p)
    assert any("nameplate MVA base is required" in str(x) or "unrated" in str(x)
               for x in cs.problems)


def test_three_winding_gets_a_control_bus_per_winding(tmp_path):
    """Run 32 failed a three-winding record with "Wind=2" until every winding
    had a control bus; run 33 loaded it once they did."""
    p = _wb(tmp_path, New_Transformers=[
        {"Scenarios": "ALL", "W1Bus": 990001, "W2Bus": 990002,
         "W3Bus": 990003, "CKT": "1", "X1_2": 0.03, "X2_3": 0.05,
         "X3_1": 0.06, "W1RateA": 350}])
    text = _render(read_workbook(p))
    header = next(l for l in text.splitlines()
                  if l.startswith("#ADD_TRANSFORMER"))
    for required in ("W3Bus#", "X2-3", "X3-1", "W1CtrlBus", "W2CtrlBus",
                     "W3CtrlBus"):
        assert required in header, f"{required} missing from {header}"


def test_two_and_three_winding_go_in_separate_sections(tmp_path):
    p = _wb(tmp_path, New_Transformers=[
        {"Scenarios": "ALL", "W1Bus": 990001, "W2Bus": 990002, "CKT": "1",
         "X1_2": 0.03, "W1RateA": 350},
        {"Scenarios": "ALL", "W1Bus": 990001, "W2Bus": 990002,
         "W3Bus": 990003, "CKT": "2", "X1_2": 0.03, "X2_3": 0.05,
         "X3_1": 0.06, "W1RateA": 350}])
    text = _render(read_workbook(p))
    assert text.count("#ADD_TRANSFORMER") == 2


def test_scale_load_uses_valueoption_not_pvalueoption(tmp_path):
    """Run 31: PValueOption is rejected and silently defaults to total Pload,
    which set area 14 load to 2 MW instead of raising it by 250 MW."""
    p = _wb(tmp_path, Gen_Load_Scaling=[
        {"Scenarios": "ALL", "Type": "LOAD", "SelectBy": "AREA",
         "Selector": "14", "ValueOption": "2", "Value": 250}])
    text = _render(read_workbook(p))
    header = next(l for l in text.splitlines() if l.startswith("#SCALE_LOAD"))
    assert "PValueOption" not in header
    assert "ValueOption" in header and "PValue" in header
