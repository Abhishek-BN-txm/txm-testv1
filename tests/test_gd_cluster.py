"""TVA GD cluster study builder + GD decks (synthetic fixtures only)."""

from __future__ import annotations

import csv
import json

import pytest

from tara_automation import gd_cluster as gc
from tara_automation.commfile import gd as gdeck
from tara_automation.footprint import Group
from tara_automation.ingest.aps import QueueData, library


# ---------------------------------------------------------------- reports ----
@pytest.mark.parametrize("expr,want", [
    ("+/- 200", (200, -200)), ("+1,000/-500", (1000, -500)), ("+588 MW/-294", (588, -294)),
    ("477/0", (477, 0)), ("+ 345/0", (345, 0)), ("345", (345, 0)), ("", None)])
def test_parse_mw(expr, want):
    assert gc.parse_mw(expr) == want


MD = """| APS Q230/SP116: +200/0 MW of Hybrid at Cholla 500 kV (200 MW SPV, 200 MW BESS)[No grid-charge] | Yes |
| APS Q275: +1,000/-500 MW of Hybrid at Cholla 500 kV (1000 MW SPV, 500 MW BESS) | Yes |
| APS Q387: +588 MW/-294 MW of Hybrid at Cholla 500 kV (588 MW SPV, 294 BESS) as ERIS | Yes |
| APS Q422: +/- 200 MW of BESS at Cholla 500 kV as NRIS | Yes |
"""


def test_parse_report_dispatch():
    d = gc.parse_report_dispatch(MD)
    assert d["Q230"]["pmax"] == 200 and d["Q230"]["pmin"] == 0 and d["Q230"]["no_grid_charge"]
    assert (d["Q275"]["pmax"], d["Q275"]["pmin"]) == (1000, -500)
    assert (d["Q387"]["pmax"], d["Q387"]["pmin"], d["Q387"]["service"]) == (588, -294, "ERIS")
    assert d["Q422"]["service"] == "NRIS" and d["Q422"]["pmin"] == -200


# --------------------------------------------------------------- topology ----
def test_radial_behind_finds_plant_side_only():
    # grid ring 1-2-...-8-1; plant 10-11-12 hangs off bus 1
    ring = list(range(1, 9))
    adj = {n: set() for n in ring + [10, 11, 12]}
    for a, b in zip(ring, ring[1:] + ring[:1]):
        adj[a].add(b)
        adj[b].add(a)
    for a, b in ((1, 10), (10, 11), (11, 12)):
        adj[a].add(b)
        adj[b].add(a)
    assert gc.radial_behind(adj, 1, max_size=5) == {10, 11, 12}
    assert gc.radial_behind(adj, 10, max_size=5) == {11, 12}
    assert gc.radial_behind(adj, 4, max_size=5) == set()


def U(bus, name, id_="1", area=14, pmin=0.0, pmax=100.0, pgen=0.0):
    return gc.Unit(bus, name, 0.6, area, id_, 1, pgen, pmin, pmax)


def test_qnum_units_matches_whole_queue_number_in_model_areas():
    units = [U(1, "Q230_SPV1"), U(2, "Q2301 BES"), U(3, "Q230 PV", area=15), U(4, "q230bes")]
    assert [u.bus for u in gc.qnum_units(units, "Q230", [14])] == [1, 4]


# ---------------------------------------------------------------- classify ---
def P(q, service="NRIS", poi="Cholla 500kV", tech="Energy Storage", rpt=None, cycle="Transition"):
    return gc.Project(q=q, cycle=cycle, mw=200, service=service, service_report=service, tech=tech,
                      status="", poi_text=poi, poi_report="", rpt=rpt or {})


def _net():
    buses = {100: library.Bus(100, "CHOLLA", 500.0, 14), 200: library.Bus(200, "FOURCORN", 345.0, 14)}
    adj = {100: {200, 101}, 200: {100}, 101: {100, 102}, 102: {101}}
    return buses, adj


def test_classify_actions():
    buses, adj = _net()
    units = [U(102, "Q230_SPV1"), U(102, "CHEV_W1A", "2")]
    ps = [P("Q230", tech="Solar"), P("Q422", rpt={"pmax": 200, "pmin": -200}),
          P("Q423", poi="Four Corners 345kV"), P("Q333", service="ERIS"),
          P("Q480", poi="Four Corners-Moenkopi 500 kV line"), P("Q273", poi="Four Corners 345kV")]
    pmap = {library.poi_key("Cholla 500kV"): {"bus": "100", "bus_name": "CHOLLA"}}
    gc.classify(ps, units, adj, buses, pmap, [14], existing={"Q273": [(102, "2")]})
    act = {p.q: p.action for p in ps}
    assert act == {"Q230": "already modelled", "Q422": "add", "Q423": "blocked", "Q333": "excluded",
                   "Q480": "blocked", "Q273": "already modelled"}
    q230 = ps[0]
    assert q230.modelled_poi == [(100, "CHOLLA")]
    q423 = ps[2]
    assert q423.candidates and q423.candidates[0][0] == 200 and q423.poi_bus is None
    # unclaimed unit behind Cholla is evidence for Q422, never a study unit
    assert any(u.name == "CHEV_W1A" for u, _ in ps[1].radial_evidence)
    assert "line-tap" in ps[4].reason


def test_new_units_pmin_and_unique_ids():
    buses, adj = _net()
    units = [U(100, "CHOLLA", "22")]   # ID '22' already used at the POI
    ps = [P("Q422", rpt={"pmax": 200, "pmin": -200}),
          P("Q424", tech="Solar", rpt={"pmax": 200, "pmin": -200}),
          P("Q230", tech="Energy Storage", rpt={"pmax": 200, "pmin": -200, "no_grid_charge": True})]
    pmap = {library.poi_key("Cholla 500kV"): {"bus": "100"}}
    gc.classify(ps, units, adj, buses, pmap, [14])
    rows = gc.new_units(ps, units)
    assert [r["GenID"] for r in rows] == ["Q2", "24", "30"]
    assert [r["GenMWMin"] for r in rows] == [-200, 0.0, 0.0]
    assert all(r["Status"] == 0 and r["GenMW"] == r["GenMWMax"] and r["BusNum"] == 100 for r in rows)


def test_change_workbook_validates(tmp_path):
    from tara_automation.casechange import read_workbook
    rows = [{"Scenarios": "ALL", "BusNum": 100, "GenID": "22", "Status": 1, "GenMW": 0.0,
             "GenMWMax": 200.0, "GenMWMin": -200.0, "Comment": "Q422"}]
    p = gc.write_change_workbook(tmp_path / "c.xlsx", rows, "GD25", "n", "d", "out.epc")
    cs = read_workbook(p)
    assert cs.ok, cs.problems
    assert [s.id for s in cs.enabled_scenarios()] == ["GD25"]
    assert cs.counts("GD25")["New_Generators"] == 1


# --------------------------------------------------------------- GI + sub ----
def test_gi_generator_rows_allocates_nris_and_labels_queue():
    units = [U(1, "Q275_SPV1", pmax=600), U(2, "Q275_BES1", pmin=-300, pmax=300),
             U(3, "OTHER", pmax=50), U(9, "FAR", area=15)]
    study = [{"q": "Q275", "bus": 1, "id": "1", "nris_mw": 600},
             {"q": "Q275", "bus": 2, "id": "1", "nris_mw": 600}]
    rows = gc.gi_generator_rows(units, [14], study, "FCC25")
    assert [r[1] for r in rows] == ["FCC25-001", "FCC25-002", "OTHER"]
    assert rows[0][3:7] == [600, 0.0, 400.0, 600] and rows[1][3:7] == [300, -300, 200.0, 300]
    assert rows[2][3:7] == [50, 0.0, 50, 50] and rows[0][8] == "Q2501"
    assert rows[0][12] == "Q2501" and rows[1][8] == "Q2501" and rows[0][9] is None


def test_write_gi_workbook_sheets(tmp_path):
    from openpyxl import load_workbook
    p = gc.write_gi_workbook(tmp_path / "gi.xlsx", [[1, "X-001", 1, 10, 0, 10, 10, None, "Q1", None, None, None, "Q1"]])
    wb = load_workbook(p)
    assert wb.sheetnames[:3] == ["Generator", "Load", "TSR"]
    assert "Line Cost" in wb.sheetnames and wb["Generator"].max_row == 2
    assert [c.value for c in wb["Generator"][1]][:7] == ["Bus#", "Bus Name", "ID", "Pmax", "Pmin", "Pmax", "Pmax"]


class _Net:
    def area_buses(self, a):
        return {2, 5, 6} if a == 14 else set()


class _FP:
    def __init__(self, groups):
        self.groups = groups
        self.network = _Net()

    def get(self, n):
        return next(g for g in self.groups if g.name == n)

    @property
    def study_area(self):
        return self.groups[0]


def test_render_gd_sub_blocks():
    fp = _FP([Group("AZ", "study_area", {1, 2}, [14]), Group("GD_SCREEN_ALL", "sub_area", {1, 2, 3}, [10, 14]),
              Group("SNK10", "sink", {3}, [10]), Group("SNK14", "sink", {2}, [])])
    text = gc.render_gd_sub(fp, sink_name="X_SINK", weights={10: 30.0, 14: 70.0}, study_name="S",
                            study_buses=[7, 7, 8], export_name="E", export_areas=[14],
                            gi_workbook="gi.xlsx", header=["h"])
    assert "SUBSYSTEM 'S'\n   BUS 7\n   BUS 8\nEND" in text
    assert "   STUDY\n      SUBSYSTEM 'S'\n   END" in text
    assert ("SUBSYSTEM 'SNK14'\n   AREA 14\n   SCALE ALL GENERATION EXCEPT\n      BUS 5\n"
            "      BUS 6\n   END\nEND") in text
    assert "SUBSYSTEM 'SNK10'\n   AREA 10\n   SCALE ALL GENERATION\nEND" in text
    assert "      SUBSYSTEM 'SNK10' 30\n      SUBSYSTEM 'SNK14' 70" in text
    assert text.rstrip().endswith('TVA_GI_DATA "gi.xlsx"')


def test_footprint_spec_excludes_region_from_export_area_sink():
    prof = gc.AREA_PROFILES["fccho"]
    spec = gc.footprint_spec(prof, "S", {10: 1.0, 14: 2.0}, [14000], "ST", exclude_buses=[5, 5, 4])
    assert "exclude" not in spec["sinks"]["SNK10"]
    ex = spec["sinks"]["SNK14"]["exclude"]
    assert ex["near"][0]["bus"] == 14000 and ex["buses"] == [4, 5]
    assert "GD_SCREEN_ALL" in spec["sub_areas"]


# ------------------------------------------------------------ load_projects --
def _w(path, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)


def test_load_projects_cycles_and_service(tmp_path):
    qd = QueueData(tmp_path)
    (tmp_path / "queue.yaml").write_text("posting: '1'\n", encoding="utf-8")
    base = {"Study Area": "Four Corners-Cholla", "Confidence": "", "Electrical Output MW": "100",
            "County (normalized)": "", "Energy Source Type": "Energy Storage", "Current Status": "",
            "Submission Year": "", "Assignment Basis": "", "Requested POI": "Cholla 500kV"}
    _w(qd.queue_by_area, [
        {**base, "FERC Q Number": "Q1", "Process": "Pre-2024 (Transition or serial)",
         "Interconnection Request Type": "Network Resource Interconnection Service"},
        {**base, "FERC Q Number": "Q2", "Process": "2025 Cluster",
         "Interconnection Request Type": "Energy Resource Interconnection Service"},
        {**base, "FERC Q Number": "Q3", "Process": "2026 Cluster",
         "Interconnection Request Type": "Network Resource Interconnection Service"},
        {**base, "FERC Q Number": "SP1", "Process": "Surplus",
         "Interconnection Request Type": "Network Resource Interconnection Service"},
        {**base, "FERC Q Number": "Q4", "Process": "2024 Cluster", "Study Area": "Gila Bend",
         "Interconnection Request Type": "Network Resource Interconnection Service"}])
    g = {"area_key": "fccho", "study_area": "", "cycle": "2025", "phase": "", "report_file": "r.pdf",
         "report_date": "2026-01-01", "is_latest_for_cycle": "True", "stage": "cluster", "q": "Q2",
         "poi": "Cholla 500 kV", "poi_kv": "500", "mw": "100", "service": "NRIS", "tech": "",
         "queue_status": "", "in_active_queue": "True", "station_nu_usd": "", "total_alloc_usd": ""}
    _w(qd.export / "generators.csv", [g])
    (qd.raw).mkdir(parents=True, exist_ok=True)
    (qd.raw / "r.pdf.json").write_text(json.dumps({"markdown": "| APS Q1: +/- 100 MW of BESS at Cholla 500 kV |"}),
                                       encoding="utf-8")
    ps = gc.load_projects(qd, "fccho", "2025")
    assert [(p.q, p.cycle, p.service) for p in ps] == [("Q1", "Serial", "NRIS"), ("Q2", "2025", "ERIS")]
    assert ps[1].service_report == "NRIS" and ps[0].rpt["pmin"] == -100
    assert [p.q for p in gc.load_projects(qd, "fccho", "2026")] == ["Q1", "Q2", "Q3"]


# ------------------------------------------------------------------ decks ----
def test_render_decks_fill_every_token():
    s = gdeck.render_screen_deck(results_dir="R", trace_log="T", case="c.epc", sub="s.sub", con="c.con",
                                 mon="m.mon", fgt="f.fgt", src="A", ref="A", scenario="sc")
    assert s.startswith("chdir") and "\nnofile.nofile\n" in s and "wdPumpMode              1" in s
    g = gdeck.render_gd_deck(results_dir="R", trace_log="T", case="c.epc", sub="s.sub", fgt="f.fgt",
                             export="E", ref="SINK", scenario="sc", ras="r.aux",
                             params={"study_queue": "FCCHO25"})
    assert "@@" not in g and g.startswith("chdir") and '"r.aux"' in g and 'StudyQueue             "FCCHO25"' in g
    assert "correctiveSwitchingMode 1" in g and "Run_TVA_GI" in g
    with pytest.raises(ValueError):
        gdeck.render_gd_deck(results_dir="R", trace_log="T", case="c", sub="s", fgt="f",
                             export="E", ref="S", scenario="x")
    with pytest.raises(KeyError):
        gdeck.render_screen_deck(results_dir="R", trace_log="T", case="c", sub="s", con="c", mon="m",
                                 fgt="f", src="A", ref="A", scenario="x", params={"bogus": 1})


def test_replay_params_match_memo():
    p = gdeck.replay_params()
    assert (p["study_selection"], p["study_queue"], p["top20_limit"], p["ignore_no_q_viol"],
            p["study_mode"], p["dc_loading_cutoff"]) == (2, "SRP26", 26, 1, 2, 90)


REPORT = """POWERENG TVA Generation Deliverability Participation Summary [x]

 Loadflow Case: c.epc
 Study Data File: s.sub
 Contingency File: not provided

Bus# , Unit Name , Gen ID , Volt , Area ,Worst Dfax
158124 , SRP26-001 , U1 , 0.48 , 15 , 0.10
158125 , SRP26-002 , U2 , 0.48 , 15 , 0.20
"""


def test_read_and_compare_reports(tmp_path):
    a, b = tmp_path / "a.csv", tmp_path / "b.csv"
    a.write_text(REPORT, encoding="utf-8")
    b.write_text(REPORT.replace("0.20", "0.205"), encoding="utf-8")
    h, rows = gdeck.read_report(a)
    assert h[0] == "Bus#" and len(rows) == 2
    assert gdeck.compare_reports(a, a)["match"]
    c = gdeck.compare_reports(a, b, tol=0.001)
    assert not c["match"] and c["cell_diffs"] == 1 and "Worst Dfax" in c["examples"][0]
    assert gdeck.compare_reports(a, b, tol=0.01)["match"]


def test_sub_helpers(tmp_path):
    s = tmp_path / "x.sub"
    s.write_text("SUBSYSTEM 'A'\nAREA 1\nEND\nSUBSYSTEM B\nEND\nTVA_GI_DATA \"gi.xlsx\"\n", encoding="utf-8")
    assert gdeck.gi_data_file(s) == tmp_path / "gi.xlsx"
    assert gdeck.defined_subsystems(s) == {"A", "B"}
