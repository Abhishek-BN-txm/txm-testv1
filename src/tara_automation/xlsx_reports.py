"""Excel workbook rendering for TARA comparisons.

Kept separate from the comparison logic so the data layer stays free of the
openpyxl dependency and remains unit-testable on its own.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F4E78")
HDR_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
TITLE = Font(name=FONT, bold=True, size=14)
SUB = Font(name=FONT, italic=True, size=9, color="555555")
BASE = Font(name=FONT, size=10)
OVR = Font(name=FONT, size=10, color="C00000", bold=True)
NEW_FILL = PatternFill("solid", fgColor="FCE4D6")
WORSE_FILL = PatternFill("solid", fgColor="FFF2CC")
RELIEF_FILL = PatternFill("solid", fgColor="E2EFDA")
_THIN = Side(style="thin", color="D0D0D0")
BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
CENTER = Alignment(horizontal="center")


def _header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)
        cell.border = BORDER


def _num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def contingency_workbook(cmp, out_path: Path) -> Path:
    """Render a ContingencyComparison to an .xlsx and return the path."""
    cases = cmp.cases
    labels = cmp.labels
    base = cases[0]
    wb = Workbook()

    # ---- Summary ----
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "N-1 Contingency Analysis — case comparison"
    ws["A1"].font = TITLE
    ws["A2"] = (f"Post-contingency rating basis: {base.rating} · baseline = "
                f"{base.label} · loading = worst post-contingency AC % per "
                f"monitored facility.")
    ws["A2"].font = SUB
    cases_line = " · ".join(
        f"{c.label} = run {c.run_number} ({c.case_name})" for c in cases)
    ws["A3"] = "Runs: " + cases_line
    ws["A3"].font = SUB

    hr = 5
    head = ["Metric"] + labels + [f"Δ {c.label} vs {base.label}"
                                  for c in cases[1:]]
    for j, h in enumerate(head, 1):
        ws.cell(row=hr, column=j, value=h)
    _header(ws, hr, len(head))

    def metric(name, vals):
        nonlocal r
        ws.cell(row=r, column=1, value=name).font = BASE
        for j, v in enumerate(vals, 2):
            ws.cell(row=r, column=j, value=v).font = BASE
            ws.cell(row=r, column=j).alignment = CENTER
        for j in range(len(vals) + 2, len(head) + 1):
            k = j - len(vals) - 2 + 1        # which post-case
            delta = vals[k] - vals[0]
            ws.cell(row=r, column=j, value=delta).font = BASE
            ws.cell(row=r, column=j).alignment = CENTER
        for c in range(1, len(head) + 1):
            ws.cell(row=r, column=c).border = BORDER
        r += 1

    r = hr + 1
    metric("Overloads ≥100% (incl. N-0 base)", [c.n_overloads for c in cases])
    metric("Solved branch records (≥cutoff)", [c.n_records for c in cases])
    metric("Non-convergent records (NConv)",
           [sum(c.nconv.values()) for c in cases])
    metric("  → distinct non-convergent contingencies",
           [len(c.nconv) for c in cases])
    metric("Voltage violations (buses)", [len(c.volt) for c in cases])

    notes = [
        "Counts use comma-safe parsing (violations.py) — the owner-name comma "
        "column-shift that inflated some loadings (e.g. GILARIVR 500 kV shown "
        "as 403–459%, truly 72–82%) is corrected here.",
        "NConv = post-contingency AC solution did not converge; see the "
        "'Non-Convergent' sheet. A change in that set between cases is a "
        "stability signal.",
    ]
    for i, t in enumerate(notes):
        ws.cell(row=r + 1 + i, column=1, value=t).font = SUB
    ws.column_dimensions["A"].width = 42
    for j in range(2, len(head) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 16

    # ---- Branch Overloads (>=100% in any case) ----
    ws2 = wb.create_sheet("Branch Overloads")
    ws2["A1"] = "Branch overloads ≥100% in any case (post-contingency)"
    ws2["A1"].font = TITLE
    ws2["A2"] = ("Colour: orange = new vs baseline, yellow = worsened, "
                 "green = relieved.")
    ws2["A2"].font = SUB
    cols = (["Monitored Facility", "kV", "Area", "Rate Base (MVA)",
             "Rate Cont (MVA)", "Contingency"]
            + [f"{lb} %" for lb in labels]
            + [f"Δ {c.label}" for c in cases[1:]] + ["Status"])
    hr2 = 4
    for j, h in enumerate(cols, 1):
        ws2.cell(row=hr2, column=j, value=h)
    _header(ws2, hr2, len(cols))
    nlab = len(labels)
    r = hr2 + 1
    for row in cmp.overload_rows:
        c = 1
        ws2.cell(row=r, column=c, value=row["facility"]); c += 1
        ws2.cell(row=r, column=c, value=row["kv"]); c += 1
        ws2.cell(row=r, column=c, value=row["area"]); c += 1
        ws2.cell(row=r, column=c, value=_num(row["rate_base"])); c += 1
        ws2.cell(row=r, column=c, value=_num(row["rate_cont"])); c += 1
        ws2.cell(row=r, column=c, value=row["contingency"]); c += 1
        load_c0 = c
        for v in row["loadings"]:
            ws2.cell(row=r, column=c, value=v); c += 1
        base_v = row["loadings"][0]
        for v in row["loadings"][1:]:
            d = round(v - base_v, 1) if (v is not None and base_v is not None) \
                else None
            ws2.cell(row=r, column=c, value=d); c += 1
        ws2.cell(row=r, column=c, value=row["status"])
        # styling
        for col in range(1, len(cols) + 1):
            cell = ws2.cell(row=r, column=col)
            cell.font = BASE
            cell.border = BORDER
            if col != 1 and col != 6 and col != len(cols):
                cell.alignment = CENTER
            if load_c0 <= col < load_c0 + nlab and isinstance(cell.value,
                                                              (int, float)):
                cell.number_format = "0.0"
                if cell.value >= 100:
                    cell.font = OVR
            if col in (4, 5) and isinstance(cell.value, (int, float)):
                cell.number_format = "0.0"
        fill = (NEW_FILL if "NEW" in row["status"]
                else WORSE_FILL if "worse" in row["status"]
                else RELIEF_FILL if "relieved" in row["status"] else None)
        if fill:
            for col in range(1, len(cols) + 1):
                ws2.cell(row=r, column=col).fill = fill
        r += 1
    ws2.freeze_panes = "A5"
    widths = [34, 8, 8, 12, 12, 34] + [9] * nlab + [8] * (nlab - 1) + [22]
    for i, w in enumerate(widths, 1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    # ---- All Monitored (>=cutoff in any case) ----
    wsa = wb.create_sheet("All Monitored")
    wsa["A1"] = "All monitored facilities ≥ cutoff in any case (post-contingency)"
    wsa["A1"].font = TITLE
    wsa["A2"] = ("Full comparison of every monitored facility+contingency that "
                 "reaches the loading cutoff in at least one case.")
    wsa["A2"].font = SUB
    colsa = ["Monitored Facility", "kV", "Area", "Rate Cont (MVA)",
             "Contingency"] + [f"{lb} %" for lb in labels]
    hra = 4
    for j, h in enumerate(colsa, 1):
        wsa.cell(row=hra, column=j, value=h)
    _header(wsa, hra, len(colsa))
    r = hra + 1
    for row in cmp.all_rows:
        vals = ([row["facility"], row["kv"], row["area"],
                 _num(row["rate_cont"]), row["contingency"]] + row["loadings"])
        for col, v in enumerate(vals, 1):
            cell = wsa.cell(row=r, column=col, value=v)
            cell.font = BASE
            cell.border = BORDER
            if col in (2, 3, 4) or col > 5:
                cell.alignment = CENTER
            if col > 5 and isinstance(v, (int, float)):
                cell.number_format = "0.0"
                if v >= 100:
                    cell.font = OVR
            if col == 4 and isinstance(v, (int, float)):
                cell.number_format = "0.0"
        r += 1
    wsa.freeze_panes = "A5"
    for i, w in enumerate([34, 8, 8, 13, 34] + [9] * nlab, 1):
        wsa.column_dimensions[get_column_letter(i)].width = w

    # ---- Voltage Violations ----
    ws3 = wb.create_sheet("Voltage Violations")
    ws3["A1"] = "Voltage violations by bus (post-contingency)"
    ws3["A1"].font = TITLE
    ws3["A2"] = "• = bus flagged in that case's N-1 voltage report."
    ws3["A2"].font = SUB
    cols3 = ["Bus #", "Bus Name", "kV", "Area"] + labels
    hr3 = 4
    for j, h in enumerate(cols3, 1):
        ws3.cell(row=hr3, column=j, value=h)
    _header(ws3, hr3, len(cols3))
    allbus = sorted({b for c in cases for b in c.volt}, key=int)
    r = hr3 + 1
    for bus in allbus:
        meta = next((c.volt[bus] for c in cases if bus in c.volt), ("", "", ""))
        ws3.cell(row=r, column=1, value=int(bus))
        ws3.cell(row=r, column=2, value=meta[0])
        ws3.cell(row=r, column=3, value=_num(meta[1]))
        ws3.cell(row=r, column=4, value=meta[2])
        for j, c in enumerate(cases, 5):
            ws3.cell(row=r, column=j, value="•" if bus in c.volt else "")
        for col in range(1, len(cols3) + 1):
            cell = ws3.cell(row=r, column=col)
            cell.font = BASE
            cell.border = BORDER
            if col != 2:
                cell.alignment = CENTER
            if col == 3 and isinstance(cell.value, (int, float)):
                cell.number_format = "0.00"
        r += 1
    ws3.freeze_panes = "A5"
    for i, w in enumerate([9, 18, 8, 10] + [7] * len(cases), 1):
        ws3.column_dimensions[get_column_letter(i)].width = w

    # ---- Non-Convergent ----
    ws4 = wb.create_sheet("Non-Convergent")
    ws4["A1"] = "Contingencies with non-convergent post-contingency AC solution"
    ws4["A1"].font = TITLE
    ws4["A2"] = ("TARA could not solve the post-contingency power flow "
                 "(loading = 'NConv'). • = non-convergent in that case.")
    ws4["A2"].font = SUB
    cols4 = ["Contingency"] + labels
    hr4 = 4
    for j, h in enumerate(cols4, 1):
        ws4.cell(row=hr4, column=j, value=h)
    _header(ws4, hr4, len(cols4))
    allc = sorted({k for c in cases for k in c.nconv})
    r = hr4 + 1
    for cont in allc:
        ws4.cell(row=r, column=1, value=cont)
        for j, c in enumerate(cases, 2):
            ws4.cell(row=r, column=j, value="•" if cont in c.nconv else "")
        for col in range(1, len(cols4) + 1):
            cell = ws4.cell(row=r, column=col)
            cell.font = BASE
            cell.border = BORDER
            if col != 1:
                cell.alignment = CENTER
        r += 1
    if not allc:
        ws4.cell(row=r, column=1,
                 value="(No non-convergent contingencies in any case.)").font = SUB
    ws4.freeze_panes = "A5"
    for i, w in enumerate([46] + [8] * len(cases), 1):
        ws4.column_dimensions[get_column_letter(i)].width = w

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path
