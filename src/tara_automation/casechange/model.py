"""Read, coerce and validate a case-change workbook into a ChangeSet.

The workbook is the user's input surface; this module turns it into data the
INCH writer can emit without further judgement. Everything that can go wrong
with a hand-filled spreadsheet is caught here and reported with sheet, row and
column, so the engineer fixes all of it in one pass instead of discovering
problems one at a time inside TARA.

Two conveniences are resolved here, not in the writer:

  * blank bus numbers on New_Buses / Line_Taps are auto-assigned from the
    scenario's NewBusStart;
  * "@Name" bus references are resolved to those numbers.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from tara_automation.casechange.schema import (
    BUS_REF_PREFIX, CHANGE_SHEETS, SCENARIOS, SHEETS, Column, Sheet,
)

EXAMPLE_MARKER = "EXAMPLE:"
ALL_SCENARIOS = "ALL"

_TRUE = {"Y", "YES", "TRUE", "T", "1"}
_FALSE = {"N", "NO", "FALSE", "F", "0"}

DEFAULT_VERSION = {"EPC": 21, "RAW": 33}


class WorkbookError(Exception):
    """Raised when the workbook cannot be read at all (missing file/sheets)."""


@dataclass(frozen=True)
class Problem:
    sheet: str
    row: int | None
    column: str | None
    message: str

    def __str__(self) -> str:
        where = self.sheet
        if self.row is not None:
            where += f" row {self.row}"
        if self.column:
            where += f" [{self.column}]"
        return f"{where}: {self.message}"


@dataclass
class Scenario:
    id: str
    name: str
    description: str = ""
    base_case: str = ""
    out_format: str = "EPC"
    out_version: int = 21
    out_name: str = ""
    solve: bool = True
    new_bus_start: int = 990001
    enabled: bool = True
    # filled by resolve(): BusName (upper) -> bus number, for @refs
    bus_names: dict[str, int] = field(default_factory=dict)

    def output_filename(self, base_case_path: Path) -> str:
        if self.out_name:
            return self.out_name
        ext = ".epc" if self.out_format == "EPC" else ".raw"
        return f"{Path(base_case_path).stem}_{self.id}{ext}"


@dataclass
class ChangeSet:
    """Every scenario plus the change rows tagged to each."""

    scenarios: list[Scenario]
    rows: dict[str, list[dict]]          # sheet name -> coerced rows
    problems: list[Problem] = field(default_factory=list)
    source: Path | None = None

    @property
    def ok(self) -> bool:
        return not self.problems

    def enabled_scenarios(self) -> list[Scenario]:
        return [s for s in self.scenarios if s.enabled]

    def scenario(self, sid: str) -> Scenario:
        return next(s for s in self.scenarios if s.id == sid)

    def rows_for(self, sheet_name: str, sid: str) -> list[dict]:
        """Rows on `sheet_name` tagged to scenario `sid` (or to ALL)."""
        out = []
        for r in self.rows.get(sheet_name, []):
            tags = r.get("Scenarios") or []
            if tags == ALL_SCENARIOS or sid in tags:
                out.append(r)
        return out

    def counts(self, sid: str) -> dict[str, int]:
        return {s.name: len(self.rows_for(s.name, sid)) for s in CHANGE_SHEETS}


# ------------------------------------------------------------- coercion ------
def _blank(v) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def _as_int(v):
    if isinstance(v, bool):
        raise ValueError("expected a number")
    if isinstance(v, int):
        return v
    if isinstance(v, float):
        if v != int(v):
            raise ValueError(f"expected a whole number, got {v}")
        return int(v)
    s = str(v).strip().replace(",", "")
    return int(float(s))


def _as_float(v):
    if isinstance(v, bool):
        raise ValueError("expected a number")
    if isinstance(v, (int, float)):
        return float(v)
    return float(str(v).strip().replace(",", ""))


def _as_flag(v) -> bool:
    if isinstance(v, bool):
        return v
    s = str(v).strip().upper()
    if s in _TRUE:
        return True
    if s in _FALSE:
        return False
    raise ValueError(f"expected Y or N, got {v!r}")


def _as_ckt(v) -> str:
    # TARA circuit IDs are 1-2 characters. Excel turns a bare 1 into 1.0, and
    # an ID of "01" into the number 1 — normalise both back to a string.
    if isinstance(v, float) and v == int(v):
        v = int(v)
    s = str(v).strip().strip("'\"")
    if not s:
        raise ValueError("circuit ID is empty")
    if len(s) > 2:
        raise ValueError(f"circuit ID must be 1-2 characters, got {s!r}")
    return s


def _as_scenarios(v):
    s = str(v).strip()
    if s.upper() == ALL_SCENARIOS:
        return ALL_SCENARIOS
    parts = [p.strip() for p in s.replace(";", ",").split(",") if p.strip()]
    if not parts:
        raise ValueError("no scenario given (use ALL or a scenario ID)")
    return parts


def _coerce(col: Column, value):
    """Coerce one cell. Returns the value, or the column default when blank."""
    if _blank(value):
        return col.default
    if col.kind == "int":
        return _as_int(value)
    if col.kind == "float":
        return _as_float(value)
    if col.kind == "flag":
        return _as_flag(value)
    if col.kind == "ckt":
        return _as_ckt(value)
    if col.kind == "scenarios":
        return _as_scenarios(value)
    if col.kind == "bus":
        s = str(value).strip()
        if s.startswith(BUS_REF_PREFIX):
            if len(s) == 1:
                raise ValueError("'@' with no bus name after it")
            return s
        return _as_int(value)
    if col.kind == "choice":
        s = str(value).strip().upper()
        if isinstance(col.default, int):        # numeric choice (ValueOption)
            s = str(_as_int(value))
        allowed = tuple(str(c).upper() for c in col.choices)
        if allowed and s not in allowed:
            raise ValueError(f"must be one of {', '.join(col.choices)}, "
                             f"got {value!r}")
        return _as_int(s) if isinstance(col.default, int) else s
    return str(value).strip()


# --------------------------------------------------------------- reading -----
def _sheet_rows(ws, sh: Sheet, problems: list[Problem]) -> list[dict]:
    """Read one worksheet into coerced dicts, skipping example rows."""
    header_cells = next(ws.iter_rows(min_row=1, max_row=1, values_only=True),
                        ())
    header = {}
    for idx, raw in enumerate(header_cells, start=1):
        if _blank(raw):
            continue
        col = sh.column(str(raw))
        if col is None:
            problems.append(Problem(sh.name, 1, str(raw),
                                    "unrecognised column — it will be ignored. "
                                    "Check the spelling against the README "
                                    "sheet."))
            continue
        header[col.name] = idx

    missing = [c.name for c in sh.columns if c.required and c.name not in header]
    if missing:
        problems.append(Problem(sh.name, 1, None,
                                "required column(s) missing from the sheet: "
                                + ", ".join(missing)))

    out: list[dict] = []
    for r, cells in enumerate(ws.iter_rows(min_row=2, values_only=True),
                              start=2):
        if all(_blank(c) for c in cells):
            continue

        def cell(name: str):
            i = header.get(name)
            return cells[i - 1] if i and i <= len(cells) else None

        # Template example rows carry the marker in their Comment cell; scan
        # every cell so a sheet without a Comment column is still covered.
        if any(isinstance(c, str) and c.strip().upper().startswith(
                EXAMPLE_MARKER) for c in cells):
            continue

        row: dict = {"_row": r}
        for col in sh.columns:
            raw = cell(col.name)
            if col.required and _blank(raw) and col.default is None:
                problems.append(Problem(sh.name, r, col.name, "required value "
                                        "is blank"))
                continue
            try:
                row[col.name] = _coerce(col, raw)
            except (ValueError, TypeError) as e:
                problems.append(Problem(sh.name, r, col.name, str(e)))
        out.append(row)
    return out


def read_workbook(path: Path) -> ChangeSet:
    """Parse a case-change workbook. Never raises on bad data — collects
    problems on the ChangeSet so the caller can report them all at once."""
    from openpyxl import load_workbook

    path = Path(path)
    if not path.exists():
        raise WorkbookError(f"Workbook not found: {path}")
    wb = load_workbook(path, data_only=True)

    problems: list[Problem] = []
    rows: dict[str, list[dict]] = {}
    for sh in SHEETS:
        if sh.name not in wb.sheetnames:
            if sh is SCENARIOS:
                raise WorkbookError(
                    f"{path.name} has no '{SCENARIOS.name}' sheet — it does not "
                    f"look like a case-change workbook. Generate a fresh one "
                    f"with: python scripts/case_change.py template")
            problems.append(Problem(sh.name, None, None,
                                    "sheet missing from the workbook — no "
                                    "changes of this type will be applied"))
            rows[sh.name] = []
            continue
        rows[sh.name] = _sheet_rows(wb[sh.name], sh, problems)

    scenarios = _build_scenarios(rows.pop(SCENARIOS.name), problems)
    cs = ChangeSet(scenarios=scenarios, rows=rows, problems=problems,
                   source=path)
    _validate(cs)
    return cs


def _build_scenarios(rows: list[dict], problems: list[Problem]) -> list[Scenario]:
    out: list[Scenario] = []
    seen: set[str] = set()
    for r in rows:
        sid = str(r.get("ScenarioID") or "").strip()
        if not sid:
            continue
        if sid.upper() == ALL_SCENARIOS:
            problems.append(Problem(SCENARIOS.name, r["_row"], "ScenarioID",
                                    "'ALL' is reserved as the tag meaning "
                                    "'every scenario' and cannot be an ID"))
            continue
        if sid in seen:
            problems.append(Problem(SCENARIOS.name, r["_row"], "ScenarioID",
                                    f"duplicate scenario ID {sid!r}"))
            continue
        seen.add(sid)
        fmt = (r.get("OutputFormat") or "EPC").upper()
        ver = r.get("OutputVersion") or DEFAULT_VERSION[fmt]
        out.append(Scenario(
            id=sid,
            name=str(r.get("Name") or sid),
            description=str(r.get("Description") or ""),
            base_case=str(r.get("BaseCase") or ""),
            out_format=fmt,
            out_version=int(ver),
            out_name=str(r.get("OutputName") or ""),
            solve=bool(r.get("Solve", True)),
            new_bus_start=int(r.get("NewBusStart") or 990001),
            enabled=bool(r.get("Enabled", True)),
        ))
    if not out:
        problems.append(Problem(SCENARIOS.name, None, None,
                                "no scenarios defined — add at least one row"))
    return out


# ------------------------------------------------------------ validation -----
def _validate(cs: ChangeSet) -> None:
    ids = {s.id for s in cs.scenarios}
    P = cs.problems

    for sh in CHANGE_SHEETS:
        for r in cs.rows.get(sh.name, []):
            tags = r.get("Scenarios")
            if tags and tags != ALL_SCENARIOS:
                for t in tags:
                    if t not in ids:
                        P.append(Problem(sh.name, r["_row"], "Scenarios",
                                         f"unknown scenario {t!r} — define it "
                                         f"on the Scenarios sheet"))

    for sc in cs.scenarios:
        _validate_scenario(cs, sc)

    for sc in cs.scenarios:
        if sc.out_format == "EPC" and sc.out_version not in (18, 19, 21):
            P.append(Problem(SCENARIOS.name, None, "OutputVersion",
                             f"scenario {sc.id}: TARA writes EPC 18, 19 or 21; "
                             f"got {sc.out_version}"))
        if sc.out_format == "RAW" and not 28 <= sc.out_version <= 33:
            P.append(Problem(SCENARIOS.name, None, "OutputVersion",
                             f"scenario {sc.id}: TARA writes RAW 28-33; got "
                             f"{sc.out_version}"))


def _validate_scenario(cs: ChangeSet, sc: Scenario) -> None:
    """Assign new bus numbers and resolve @refs for one scenario."""
    P = cs.problems
    next_num = sc.new_bus_start
    taken: set[int] = set()
    names: dict[str, int] = {}

    def claim(num: int | None, name: str, sheet_name: str, row: int,
              col: str) -> int | None:
        nonlocal next_num
        if num is None:
            while next_num in taken:
                next_num += 1
            num = next_num
            next_num += 1
        if num in taken:
            P.append(Problem(sheet_name, row, col,
                             f"bus number {num} is used twice in scenario "
                             f"{sc.id}"))
            return num
        taken.add(num)
        key = name.strip().upper()
        if key in names:
            P.append(Problem(sheet_name, row, "BusName",
                             f"bus name {name!r} is used twice in scenario "
                             f"{sc.id}; @-references would be ambiguous"))
        else:
            names[key] = num
        return num

    for r in cs.rows_for("New_Buses", sc.id):
        r.setdefault("_num", {})[sc.id] = claim(
            r.get("BusNum"), str(r.get("BusName") or ""), "New_Buses",
            r["_row"], "BusNum")
    for r in cs.rows_for("Line_Taps", sc.id):
        r.setdefault("_num", {})[sc.id] = claim(
            r.get("NewBusNum"), str(r.get("NewBusName") or ""), "Line_Taps",
            r["_row"], "NewBusNum")

    sc.bus_names = names

    # Resolve @refs on every bus-typed column.
    for sh in CHANGE_SHEETS:
        bus_cols = [c.name for c in sh.columns if c.kind == "bus"]
        if not bus_cols:
            continue
        for r in cs.rows_for(sh.name, sc.id):
            for cname in bus_cols:
                v = r.get(cname)
                if not isinstance(v, str) or not v.startswith(BUS_REF_PREFIX):
                    continue
                key = v[1:].strip().upper()
                if key in names:
                    r.setdefault("_resolved", {}).setdefault(sc.id, {})[
                        cname] = names[key]
                else:
                    P.append(Problem(
                        sh.name, r["_row"], cname,
                        f"{v!r} does not match any BusName on New_Buses or "
                        f"NewBusName on Line_Taps for scenario {sc.id}"))

    # Bus name length: EPC/RAW carry 12 characters.
    for r in cs.rows_for("New_Buses", sc.id):
        nm = str(r.get("BusName") or "")
        if len(nm) > 12:
            P.append(Problem("New_Buses", r["_row"], "BusName",
                             f"{nm!r} is {len(nm)} characters; EPC/RAW hold 12 "
                             f"— TARA will truncate it"))

    # A new line with no rating is invisible to monitoring; catch it here.
    for r in cs.rows_for("New_Branches", sc.id):
        if not r.get("RateA"):
            P.append(Problem("New_Branches", r["_row"], "RateA",
                             "rating is zero or blank — TARA treats an unrated "
                             "branch as unmonitorable, so it will never show a "
                             "violation"))

    for r in cs.rows_for("New_Transformers", sc.id):
        if not r.get("W1RateA"):
            P.append(Problem("New_Transformers", r["_row"], "W1RateA",
                             "rating is zero or blank — an unrated transformer "
                             "is never monitored"))
        # NAMEPLATE_PCT divides by the nameplate base; without one the
        # conversion is meaningless and would silently pass the percentage
        # through as system-base per unit — a 350x error.
        if (r.get("ImpedanceBasis") or "") == "NAMEPLATE_PCT":
            base = r.get("SBAS1") or r.get("W1RateA")
            if not base:
                P.append(Problem(
                    "New_Transformers", r["_row"], "SBAS1",
                    "ImpedanceBasis is NAMEPLATE_PCT, so a nameplate MVA base "
                    "is required — set SBAS1 (or W1RateA)"))
        if cs.rows and r.get("X2_3") and not r.get("W3Bus"):
            P.append(Problem(
                "New_Transformers", r["_row"], "W3Bus",
                "2-3 winding impedance given but no W3Bus — set W3Bus for a "
                "three-winding transformer, or clear the 2-3 / 3-1 columns"))


def bus_value(row: dict, column: str, scenario_id: str):
    """The resolved bus number for a bus-typed cell (handles @refs)."""
    resolved = row.get("_resolved", {}).get(scenario_id, {})
    if column in resolved:
        return resolved[column]
    return row.get(column)


def new_bus_number(row: dict, scenario_id: str):
    """The assigned number for a New_Buses / Line_Taps row."""
    return row.get("_num", {}).get(scenario_id)
